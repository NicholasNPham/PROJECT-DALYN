"""Check every rule's Type/Subtype against STAC's own list, before STAC sees it.

    python check_types.py

Reads both spreadsheets named in config.yaml:
    paths.excel       the rules sheet, e.g. DALYN26.xlsx
    paths.stac_types  STAC's export of Type/Subtype pairs

Every rule that files to a pair STAC does not have is printed with the
closest real pairs, so the sheet can be fixed in one pass rather than one
failed document at a time inside STAC.

Read-only. Touches no mailbox, opens no browser, changes no spreadsheet.
Exit code 0 when every rule is good, 1 when any rule is not.
"""

import sys
from difflib import get_close_matches
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent / "src"))
import classifier  # noqa: E402
from config_loader import (  # noqa: E402
    LIVE_REVIEW_SUBTYPE,
    LIVE_REVIEW_TYPE,
    load_config,
)
from exceptions import SystemProblem  # noqa: E402
from logger import setup_logging  # noqa: E402

# Columns of STAC's export, by header name rather than position, since the
# export is somebody else's file and its column order is not ours to rely on.
TYPE_HEADER = "type"
SUBTYPE_HEADER = "subtype"
DESCRIPTION_HEADER = "description"
INACTIVE_HEADER = "inactive"

SUGGESTIONS = 3


def load_stac_pairs(path: Path) -> dict[tuple[str, str], tuple[str, bool]]:
    """Return {(TYPE, SUBTYPE): (description, active)} from STAC's export."""
    worksheet = load_workbook(path, data_only=True, read_only=True).active
    rows = worksheet.iter_rows(values_only=True)

    try:
        header = [str(cell).strip().lower() if cell else "" for cell in next(rows)]
    except StopIteration:
        raise SystemProblem(f"{path} is empty.") from None

    try:
        type_at = header.index(TYPE_HEADER)
        subtype_at = header.index(SUBTYPE_HEADER)
    except ValueError:
        raise SystemProblem(
            f"{path} has no '{TYPE_HEADER}' and '{SUBTYPE_HEADER}' columns. "
            f"Found: {', '.join(h for h in header if h)}"
        ) from None

    description_at = header.index(DESCRIPTION_HEADER) if DESCRIPTION_HEADER in header else None
    inactive_at = header.index(INACTIVE_HEADER) if INACTIVE_HEADER in header else None

    pairs = {}
    for row in rows:
        if not row or not row[type_at] or not row[subtype_at]:
            continue
        key = (str(row[type_at]).strip().upper(), str(row[subtype_at]).strip().upper())
        description = str(row[description_at]).strip() if description_at is not None and row[description_at] else ""
        # Anything other than a clear Y counts as active, which is the
        # forgiving direction: a pair wrongly called active fails later in
        # STAC, where a pair wrongly called inactive would send us chasing
        # a rule that was fine.
        inactive = inactive_at is not None and str(row[inactive_at]).strip().upper() == "Y"
        pairs[key] = (description if description.upper() != "NULL" else "", not inactive)

    return pairs


def check_review_pair(
    config: dict,
    pairs: dict,
    active: set,
    flat_active: list[str],
) -> bool:
    """Check the pair unclassified attachments are filed under. True if broken.

    This pair does not come from the rules sheet, so the loop over rules never
    saw it. That gap is exactly how a run reached STAC and failed on every
    unclassified document with "no row for PLS/RVW": the rules all checked out
    and the one hardcoded pair was never looked at.

    The fallbacks come from config_loader rather than being written out here.
    They should never fire, since _validate_review_pair fills both keys, but a
    literal "PLS"/"RVW" written in this file would be the exact wrong pair
    sitting inside the script whose job is catching wrong pairs.
    """
    stac = config.get("stac", {})
    key = (
        str(stac.get("review_type", LIVE_REVIEW_TYPE)).strip().upper(),
        str(stac.get("review_subtype", LIVE_REVIEW_SUBTYPE)).strip().upper(),
    )
    label = f"{key[0]}/{key[1]}"

    if key in active:
        print(f"Unclassified attachments file to {label}, which is active in STAC.\n")
        return False

    if key in pairs:
        print(f"PROBLEM: unclassified attachments file to {label}, INACTIVE in STAC.")
    else:
        print(f"PROBLEM: unclassified attachments file to {label}, NOT in STAC.")
        for suggestion in get_close_matches(label, flat_active, n=SUGGESTIONS, cutoff=0.3):
            document_type, subtype = suggestion.split("/", 1)
            description = pairs[(document_type, subtype)][0]
            print(f"         try {suggestion}{'  ' + description if description else ''}")

    print(
        "         Set stac.review_type and stac.review_subtype in config.yaml to a\n"
        "         pair this instance has. Every attachment that matches no rule\n"
        "         fails in STAC until then.\n"
    )
    return True


def main() -> int:
    try:
        config = load_config(with_credentials=False)
    except SystemProblem as error:
        print(f"Config problem: {error}", file=sys.stderr)
        return 1

    setup_logging(config["paths"]["logs"])

    rules_path = config["paths"]["excel"]
    types_path = config["paths"]["stac_types"]

    try:
        pairs = load_stac_pairs(types_path)
        rules = classifier.load_rules(rules_path)
    except SystemProblem as error:
        print(f"Could not read a spreadsheet: {error}", file=sys.stderr)
        return 1

    active = {key for key, (_, is_active) in pairs.items() if is_active}
    flat_active = sorted(f"{t}/{st}" for t, st in active)

    print(f"\nRules sheet: {rules_path}")
    print(f"STAC types:  {types_path}")
    print(f"{len(rules)} rules, {len(pairs)} pairs in STAC ({len(active)} active)\n")

    problems = []
    for rule in rules:
        # A blank Type and Subtype is a deliberate safety net, not a target.
        if not rule.document_type and not rule.document_subtype:
            continue

        key = (
            (rule.document_type or "").strip().upper(),
            (rule.document_subtype or "").strip().upper(),
        )
        if key in active:
            continue

        if key in pairs:
            problems.append((rule, key, "INACTIVE in STAC", []))
            continue

        suggestions = get_close_matches(f"{key[0]}/{key[1]}", flat_active, n=SUGGESTIONS, cutoff=0.5)
        problems.append((rule, key, "NOT in STAC", suggestions))

    review_problem = check_review_pair(config, pairs, active, flat_active)

    if not problems and not review_problem:
        print("Every rule files to an active STAC Type/Subtype pair.\n")
        return 0

    if not problems:
        return 1

    print(f"{len(problems)} of {len(rules)} rules point at a pair STAC will not accept:\n")
    for rule, key, what, suggestions in problems:
        print(f"  row {rule.row:>3}  {rule.phrase}")
        print(f"           files to {key[0]}/{key[1]}  <-- {what}")
        for suggestion in suggestions:
            t, st = suggestion.split("/", 1)
            description = pairs[(t, st)][0]
            print(f"           try      {suggestion}{'  ' + description if description else ''}")
        print()

    print("Fix these in the rules sheet, then run this again.\n")
    return 1


if __name__ == "__main__":
    sys.exit(main())